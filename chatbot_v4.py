
"""
Master Chief - Tutor Virtual
Version 4.0

Ahora puede trabajar con documentos, imágenes y audio.
"""

import os
import json
import uuid
import base64
from html import escape
from pathlib import Path

import streamlit as st
from dotenv import load_dotenv

from openai import OpenAI
from langchain_openai import ChatOpenAI
from langchain_core.messages import SystemMessage, HumanMessage, AIMessage

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from utils import cargar_estilos


# Configuración inicial

st.set_page_config(
    page_title="Master Chief - Tutor Virtual",
    page_icon="🤖"
)

load_dotenv()

API_KEY = os.getenv("OPENAI_API_KEY")

if not API_KEY:
    st.error("No se encontró OPENAI_API_KEY en el archivo .env.")
    st.stop()

cliente_openai = OpenAI(api_key=API_KEY)

cargar_estilos("styles.css")


MODELOS_DISPONIBLES = {
    "GPT-4o mini (rápido y económico)": "gpt-4o-mini",
    "GPT-4o (más capaz, más costoso)": "gpt-4o",
}


SYSTEM_PROMPT_BASE = (
    "Eres Master Chief, un asistente de inteligencia artificial "
    "que ayuda al estudiante de manera cercana, natural y clara. "
    "Si te preguntan si eres una IA, responde con honestidad."
)


INSTRUCCIONES_MODO = {
    "📚 Tutor": (
        "Actúa como un tutor académico. Explica paso a paso, "
        "utiliza ejemplos sencillos y ayuda al estudiante "
        "a comprender, no solamente a memorizar respuestas."
    ),

    "📝 Práctica": (
        "Actúa como un profesor que prepara al estudiante. "
        "Genera preguntas, permite que responda y después "
        "explica sus aciertos y errores."
    ),

    "💡 Explicación rápida": (
        "Explica de manera breve, directa y sencilla, "
        "sin omitir información importante."
    ),
}


INSTRUCCIONES_NIVEL = {
    "🎓 Básico": (
        "Utiliza lenguaje sencillo, define términos técnicos "
        "y explica con ejemplos fáciles."
    ),

    "📘 Intermedio": (
        "Utiliza conceptos técnicos cuando sean necesarios, "
        "explicando los más importantes."
    ),

    "📚 Universitario": (
        "Utiliza terminología académica y desarrolla los "
        "conceptos con suficiente profundidad."
    ),

    "🧠 Avanzado": (
        "Profundiza en los conceptos y utiliza terminología "
        "técnica especializada cuando corresponda."
    ),
}


# Archivos donde se guardan las conversaciones

DATA_DIR = Path("data")
DATA_DIR.mkdir(exist_ok=True)

CONVERSACIONES_PATH = DATA_DIR / "conversaciones.json"
HISTORIAL_ANTERIOR_PATH = DATA_DIR / "historial.json"


# Conversaciones

def crear_conversacion(titulo="Nueva conversación"):
    return {
        "id": uuid.uuid4().hex[:8],
        "titulo": titulo,
        "mensajes": []
    }


def buscar_conversaciones(conversaciones, consulta):
    consulta = consulta.strip().casefold()

    if not consulta:
        return conversaciones

    resultados = []

    for conversacion in conversaciones:
        contenido = " ".join(
            str(mensaje.get("content", ""))
            for mensaje in conversacion.get("mensajes", [])
        )
        texto_busqueda = (
            str(conversacion.get("titulo", "")) + " " + contenido
        ).casefold()

        if consulta in texto_busqueda:
            resultados.append(conversacion)

    return resultados


def obtener_fragmento_coincidente(conversacion, consulta):
    consulta = consulta.strip()

    if not consulta:
        return ""

    consulta_normalizada = consulta.casefold()

    for mensaje in conversacion.get("mensajes", []):
        contenido = str(mensaje.get("content", "")).replace("\n", " ")
        indice = contenido.casefold().find(consulta_normalizada)

        if indice >= 0:
            inicio = max(0, indice - 35)
            fin = min(len(contenido), indice + len(consulta) + 55)
            prefijo = "..." if inicio else ""
            sufijo = "..." if fin < len(contenido) else ""
            rol = "Tú" if mensaje.get("role") == "user" else "Master Chief"
            return f"{rol}: {prefijo}{contenido[inicio:fin]}{sufijo}"

    return ""


def guardar_conversaciones(conversaciones):
    CONVERSACIONES_PATH.write_text(
        json.dumps(conversaciones, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )


def cargar_conversaciones():
    if CONVERSACIONES_PATH.exists():
        try:
            conversaciones = json.loads(
                CONVERSACIONES_PATH.read_text(encoding="utf-8")
            )

            if isinstance(conversaciones, list) and conversaciones:
                return conversaciones

        except (json.JSONDecodeError, OSError):
            pass

    if HISTORIAL_ANTERIOR_PATH.exists():
        try:
            historial = json.loads(
                HISTORIAL_ANTERIOR_PATH.read_text(encoding="utf-8")
            )

            if isinstance(historial, list) and historial:
                conversacion = crear_conversacion("Conversación anterior")
                conversacion["mensajes"] = historial

                conversaciones = [conversacion]
                guardar_conversaciones(conversaciones)

                return conversaciones

        except (json.JSONDecodeError, OSError):
            pass

    conversaciones = [crear_conversacion()]
    guardar_conversaciones(conversaciones)

    return conversaciones


if "conversaciones" not in st.session_state:
    st.session_state.conversaciones = cargar_conversaciones()

if "chat_activo" not in st.session_state:
    st.session_state.chat_activo = (
        st.session_state.conversaciones[0]["id"]
    )

CLAVES_CONTEXTO_ARCHIVOS = (
    "documento_fragmentos",
    "documento_nombres",
    "imagenes_activas",
    "audio_transcripciones",
    "audio_nombres"
)

if "contextos_archivos" not in st.session_state:
    contexto_heredado = {
        clave: st.session_state.pop(clave)
        for clave in CLAVES_CONTEXTO_ARCHIVOS
        if clave in st.session_state
    }
    st.session_state.contextos_archivos = {}

    if contexto_heredado:
        st.session_state.contextos_archivos[
            st.session_state.chat_activo
        ] = contexto_heredado


def obtener_contexto_archivos(conversacion_id):
    return st.session_state.contextos_archivos.setdefault(
        conversacion_id,
        {}
    )


def obtener_conversacion_activa():
    for conversacion in st.session_state.conversaciones:
        if conversacion["id"] == st.session_state.chat_activo:
            return conversacion

    nueva = crear_conversacion()
    st.session_state.conversaciones.insert(0, nueva)
    st.session_state.chat_activo = nueva["id"]

    guardar_conversaciones(st.session_state.conversaciones)

    return nueva


conversacion_actual = obtener_conversacion_activa()
st.session_state.messages = conversacion_actual["mensajes"]


# Lectura de documentos

def leer_pdf(archivo):
    lector = PdfReader(archivo)

    texto = []

    for pagina in lector.pages:
        contenido = pagina.extract_text()

        if contenido:
            texto.append(contenido)

    return "\n".join(texto)


def leer_txt(archivo):
    contenido = archivo.read()

    if isinstance(contenido, bytes):
        contenido = contenido.decode("utf-8", errors="ignore")

    return contenido


def leer_docx(archivo):
    documento = Document(archivo)
    texto = []

    for parrafo in documento.paragraphs:
        if parrafo.text.strip():
            texto.append(parrafo.text)

    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [celda.text.strip() for celda in fila.cells]
            texto.append(" | ".join(celdas))

    return "\n".join(texto)


def leer_xlsx(archivo):
    libro = load_workbook(
        archivo,
        read_only=True,
        data_only=True
    )

    texto = []

    for hoja in libro.worksheets:
        texto.append(f"\n--- Hoja: {hoja.title} ---\n")

        for fila in hoja.iter_rows(values_only=True):
            valores = [
                str(valor)
                for valor in fila
                if valor is not None
            ]

            if valores:
                texto.append(" | ".join(valores))

    return "\n".join(texto)


def leer_archivo(archivo):
    nombre = archivo.name.lower()

    if nombre.endswith(".pdf"):
        return leer_pdf(archivo)

    if nombre.endswith(".txt"):
        return leer_txt(archivo)

    if nombre.endswith(".docx"):
        return leer_docx(archivo)

    if nombre.endswith(".xlsx"):
        return leer_xlsx(archivo)

    return ""


# Dividir documentos para encontrar información útil

def dividir_texto(texto, tamano=1000):
    palabras = texto.split()
    fragmentos = []

    for inicio in range(0, len(palabras), tamano):
        fragmento = " ".join(palabras[inicio:inicio + tamano])

        if fragmento.strip():
            fragmentos.append(fragmento)

    return fragmentos


def buscar_informacion_relevante(pregunta, fragmentos, cantidad=3):
    if not fragmentos:
        return []

    vectorizador = TfidfVectorizer()
    matriz = vectorizador.fit_transform(fragmentos + [pregunta])

    similitudes = cosine_similarity(
        matriz[-1],
        matriz[:-1]
    )[0]

    indices = similitudes.argsort()[::-1][:cantidad]

    return [
        fragmentos[i]
        for i in indices
        if similitudes[i] > 0
    ]


# Transcribir audio

def transcribir_audio(archivo):
    extensiones_validas = {
        "mp3", "mp4", "mpeg", "mpga",
        "m4a", "wav", "webm"
    }

    extension = archivo.name.rsplit(".", 1)[-1].lower()

    if extension not in extensiones_validas:
        raise ValueError("El formato de audio no es compatible.")

    archivo.seek(0)

    resultado = cliente_openai.audio.transcriptions.create(
        model="gpt-4o-mini-transcribe",
        file=(
            archivo.name,
            archivo.getvalue(),
            archivo.type or "application/octet-stream"
        )
    )

    return resultado.text


def procesar_archivos_adjuntos(archivos, conversacion_id):
    extensiones_documento = {"pdf", "txt", "docx", "xlsx"}
    extensiones_imagen = {"jpg", "jpeg", "png", "webp"}
    extensiones_audio = {"mp3", "mp4", "mpeg", "mpga", "m4a", "wav", "webm"}

    fragmentos_documento = []
    nombres_documento = []
    imagenes = []
    transcripciones_audio = []
    nombres_audio = []

    for archivo in archivos:
        extension = archivo.name.rsplit(".", 1)[-1].lower()

        try:
            if extension in extensiones_documento:
                texto = leer_archivo(archivo)

                if texto.strip():
                    fragmentos_documento.extend(
                        dividir_texto(f"Archivo: {archivo.name}\n{texto}")
                    )
                    nombres_documento.append(archivo.name)
                else:
                    st.warning(f"No se encontró texto en {archivo.name}.")

            elif extension in extensiones_imagen:
                imagenes.append({
                    "nombre": archivo.name,
                    "bytes": archivo.getvalue(),
                    "tipo": archivo.type or "image/png"
                })

            elif extension in extensiones_audio:
                with st.spinner(f"Transcribiendo {archivo.name}..."):
                    transcripciones_audio.append(
                        f"Archivo: {archivo.name}\n{transcribir_audio(archivo)}"
                    )
                nombres_audio.append(archivo.name)

        except Exception as error:
            st.error(f"No se pudo procesar {archivo.name}: {error}")

    contexto = obtener_contexto_archivos(conversacion_id)

    if fragmentos_documento:
        contexto["documento_fragmentos"] = fragmentos_documento
        contexto["documento_nombres"] = nombres_documento

    if imagenes:
        contexto["imagenes_activas"] = imagenes

    if transcripciones_audio:
        contexto["audio_transcripciones"] = "\n\n".join(
            transcripciones_audio
        )
        contexto["audio_nombres"] = nombres_audio


# Título

st.title("Master Chief, tu tutor virtual")

st.caption(
    "Versión 4.0 — documentos, imágenes y audio"
)


# Panel lateral

with st.sidebar:

    st.header("Chats")

    consulta_chats = st.text_input(
        "Buscar conversaciones",
        placeholder="Buscar por tema o palabra...",
        key="buscar_conversaciones"
    )

    if st.button("➕ Nueva conversación", use_container_width=True):

        nueva = crear_conversacion()

        st.session_state.conversaciones.insert(0, nueva)
        st.session_state.chat_activo = nueva["id"]

        guardar_conversaciones(st.session_state.conversaciones)

        st.rerun()

    chats_visibles = buscar_conversaciones(
        st.session_state.conversaciones,
        consulta_chats
    )

    st.caption(f"{len(chats_visibles)} resultados")

    for conversacion in chats_visibles:
        with st.container(key=f"fila_chat_{conversacion['id']}"):
            columna_chat, columna_eliminar = st.columns(
                [0.84, 0.16],
                gap="small"
            )

            with columna_chat:
                if st.button(
                    conversacion["titulo"],
                    key=f"chat_{conversacion['id']}",
                    type=(
                        "primary"
                        if conversacion["id"] == st.session_state.chat_activo
                        else "secondary"
                    ),
                    use_container_width=True
                ):
                    st.session_state.chat_activo = conversacion["id"]
                    st.rerun()

            with columna_eliminar:
                if st.button(
                    " ",
                    icon=":material/delete:",
                    type="tertiary",
                    key=f"eliminar_chat_{conversacion['id']}",
                    help=f"Eliminar {conversacion['titulo']}"
                ):
                    st.session_state.chat_eliminar_pendiente = conversacion["id"]
                    st.rerun()

            coincidencia = obtener_fragmento_coincidente(
                conversacion,
                consulta_chats
            )
            if coincidencia:
                st.caption(coincidencia)

    if not chats_visibles:
        st.caption("No hay conversaciones que coincidan con la búsqueda.")

    conversacion_pendiente_id = st.session_state.get(
        "chat_eliminar_pendiente"
    )

    if conversacion_pendiente_id:
        conversacion_pendiente = next(
            (
                conversacion
                for conversacion in st.session_state.conversaciones
                if conversacion["id"] == conversacion_pendiente_id
            ),
            None
        )

        if conversacion_pendiente:
            st.warning(
                f"¿Eliminar '{conversacion_pendiente['titulo']}'? "
                "Esta acción no se puede deshacer."
            )
            columna_confirmar, columna_cancelar = st.columns(2)

            with columna_confirmar:
                if st.button(
                    "Eliminar",
                    key="confirmar_eliminar_chat",
                    type="primary",
                    use_container_width=True
                ):
                    st.session_state.conversaciones = [
                        conversacion
                        for conversacion in st.session_state.conversaciones
                        if conversacion["id"] != conversacion_pendiente_id
                    ]
                    st.session_state.contextos_archivos.pop(
                        conversacion_pendiente_id,
                        None
                    )

                    if not st.session_state.conversaciones:
                        st.session_state.conversaciones.append(
                            crear_conversacion()
                        )

                    if st.session_state.chat_activo == conversacion_pendiente_id:
                        st.session_state.chat_activo = (
                            st.session_state.conversaciones[0]["id"]
                        )

                    st.session_state.pop("chat_eliminar_pendiente", None)
                    guardar_conversaciones(st.session_state.conversaciones)
                    st.rerun()

            with columna_cancelar:
                if st.button(
                    "Cancelar",
                    key="cancelar_eliminar_chat",
                    use_container_width=True
                ):
                    st.session_state.pop("chat_eliminar_pendiente", None)
                    st.rerun()

    st.divider()

    # Configuración

    st.header("⚙️ Configuración")

    modo_aprendizaje = st.selectbox(
        "Modo de aprendizaje",
        list(INSTRUCCIONES_MODO.keys())
    )

    nivel_academico = st.selectbox(
        "Nivel académico",
        list(INSTRUCCIONES_NIVEL.keys()),
        index=2
    )

    modelo_visible = st.selectbox(
        "Modelo de OpenAI",
        list(MODELOS_DISPONIBLES.keys()),
        index=0
    )

    MODEL_NAME = MODELOS_DISPONIBLES[modelo_visible]

    max_historial = st.slider(
        "Mensajes de historial a recordar",
        min_value=2,
        max_value=30,
        value=10,
        step=2
    )

    if st.button("🗑️ Limpiar conversación", use_container_width=True):

        conversacion_actual["mensajes"] = []
        conversacion_actual["titulo"] = "Nueva conversación"

        guardar_conversaciones(st.session_state.conversaciones)

        st.rerun()

    if st.session_state.messages:

        texto_exportado = "\n\n".join(
            f"{'Tú' if m['role'] == 'user' else 'Master Chief'}: "
            f"{m['content']}"
            for m in st.session_state.messages
        )

        st.download_button(
            "⬇️ Descargar conversación",
            data=texto_exportado,
            file_name="conversacion_master_chief.txt",
            mime="text/plain",
            use_container_width=True
        )

    st.divider()

    st.caption(
        "Master Chief es un asistente de IA con fines académicos. "
        "Las conversaciones se guardan localmente. "
        "Las respuestas pueden contener errores."
    )


# Modelo de respuesta

llm = ChatOpenAI(
    model=MODEL_NAME,
    temperature=0,
    api_key=API_KEY
)


entrada = st.chat_input(
    "Escribe tu mensaje...",
    accept_file="multiple",
    file_type=[
        "pdf", "txt", "docx", "xlsx",
        "jpg", "jpeg", "png", "webp",
        "mp3", "mp4", "mpeg", "mpga", "m4a", "wav", "webm"
    ]
)

if entrada:
    prompt = entrada.text.strip()
    archivos_adjuntos = entrada.files

    if archivos_adjuntos:
        procesar_archivos_adjuntos(
            archivos_adjuntos,
            conversacion_actual["id"]
        )

    if not prompt and archivos_adjuntos:
        prompt = "Analiza los archivos adjuntos y resume la información más importante."

    if not prompt:
        st.stop()


# Mostrar mensajes anteriores

for mensaje in st.session_state.messages:
    with st.chat_message(mensaje["role"]):
        st.markdown(mensaje["content"])


nombres_archivos_activos = (
    obtener_contexto_archivos(conversacion_actual["id"]).get(
        "documento_nombres",
        []
    )
    + obtener_contexto_archivos(conversacion_actual["id"]).get(
        "audio_nombres",
        []
    )
    + [
        imagen["nombre"]
        for imagen in obtener_contexto_archivos(
            conversacion_actual["id"]
        ).get("imagenes_activas", [])
    ]
)

if nombres_archivos_activos:
    estado_archivos, boton_quitar = st.columns([5, 1])
    with estado_archivos:
        archivos_html = "".join(
            f'<span class="file-context-item">{escape(nombre)}</span>'
            for nombre in nombres_archivos_activos
        )
        st.markdown(
            '<div class="file-context" role="group" '
            'aria-label="Archivos de contexto activos">'
            '<span class="file-context-label">Contexto activo</span>'
            f"{archivos_html}</div>",
            unsafe_allow_html=True
        )

    with boton_quitar:
        if st.button("Quitar", key="quitar_contexto_archivos"):
            st.session_state.contextos_archivos.pop(
                conversacion_actual["id"],
                None
            )

            st.rerun()


# Chat principal
if entrada:
    if archivos_adjuntos:
        prompt_visible = (
            prompt
            + "\n\n📎 "
            + ", ".join(archivo.name for archivo in archivos_adjuntos)
        )
    else:
        prompt_visible = prompt

    st.chat_message("user").markdown(prompt_visible)

    st.session_state.messages.append({
        "role": "user",
        "content": prompt_visible
    })

    # El primer mensaje sirve como título de la conversación

    if conversacion_actual["titulo"] == "Nueva conversación":

        titulo = (
            entrada.text.strip()
            or ", ".join(archivo.name for archivo in archivos_adjuntos)
        ).replace("\n", " ")

        if len(titulo) > 35:
            titulo = titulo[:35] + "..."

        conversacion_actual["titulo"] = titulo

    st.session_state.conversaciones.remove(conversacion_actual)
    st.session_state.conversaciones.insert(0, conversacion_actual)

    guardar_conversaciones(st.session_state.conversaciones)

    historial = st.session_state.messages[-max_historial:]

    # Agregar el contenido relacionado con los archivos

    contexto = ""

    contexto_archivos = obtener_contexto_archivos(
        conversacion_actual["id"]
    )

    if contexto_archivos.get("documento_fragmentos"):

        relevantes = buscar_informacion_relevante(
            prompt,
            contexto_archivos["documento_fragmentos"],
            cantidad=3
        )

        if relevantes:
            contexto += (
                "\n\nInformación relevante del documento:\n"
                + "\n\n".join(relevantes)
                + "\nPrioriza esta información cuando corresponda. "
                "No inventes datos que no aparezcan en ella."
            )

    if contexto_archivos.get("audio_transcripciones"):
        contexto += (
            "\n\nTranscripción del audio:\n"
            + contexto_archivos["audio_transcripciones"]
            + "\nUtiliza esta transcripción para responder."
        )

    system_prompt = (
        SYSTEM_PROMPT_BASE
        + "\n\n"
        + INSTRUCCIONES_MODO[modo_aprendizaje]
        + "\n\n"
        + INSTRUCCIONES_NIVEL[nivel_academico]
        + contexto
    )

    lc_messages = [
        SystemMessage(content=system_prompt)
    ]

    for mensaje in historial[:-1]:

        if mensaje["role"] == "user":
            lc_messages.append(HumanMessage(content=mensaje["content"]))

        else:
            lc_messages.append(AIMessage(content=mensaje["content"]))

    # Si hay una imagen, se envía junto con la pregunta

    contenido_usuario = [
        {
            "type": "text",
            "text": prompt
        }
    ]

    for imagen in contexto_archivos.get("imagenes_activas", []):
        imagen_base64 = base64.b64encode(
            imagen["bytes"]
        ).decode("utf-8")

        imagen_url = (
            f"data:{imagen['tipo']};base64,"
            f"{imagen_base64}"
        )

        contenido_usuario.append({
            "type": "image_url",
            "image_url": {
                "url": imagen_url
            }
        })

    lc_messages.append(
        HumanMessage(content=contenido_usuario)
    )

    # Mostrar la respuesta poco a poco

    with st.chat_message("assistant"):

        try:

            def generar_respuesta():
                for chunk in llm.stream(lc_messages):
                    if isinstance(chunk.content, str) and chunk.content:
                        yield chunk.content

            respuesta = st.write_stream(generar_respuesta)

        except Exception as e:

            respuesta = f"Ocurrió un error al generar la respuesta: {e}"
            st.error(respuesta)

    st.session_state.messages.append({
        "role": "assistant",
        "content": respuesta
    })

    guardar_conversaciones(st.session_state.conversaciones)