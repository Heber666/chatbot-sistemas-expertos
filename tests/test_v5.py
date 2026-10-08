"""Pruebas locales de v5; no usan datos personales ni solicitudes reales."""

import base64
import json
import os
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import httpx
from docx import Document
from PIL import Image, ImageDraw
from langchain_core.messages import AIMessageChunk
from openai import APIConnectionError, AuthenticationError, NotFoundError, RateLimitError
from openpyxl import Workbook
from pydantic import ValidationError
from pypdf import PdfReader
from streamlit.testing.v1 import AppTest

import archivos
import storage
from busqueda import buscar_informacion_relevante, dividir_texto
from quiz import Quiz, calificar_quiz, generar_quiz

ROOT = Path(__file__).resolve().parents[1]


def pdf_escaneado(cantidad=1):
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    imagen = Image.new("RGB", (600, 600), "white")
    ImageDraw.Draw(imagen).text((30, 30), "CODIGO DE ETICA\nTrabaja con honestidad.\nRespeta a tus companeros.", fill="black")
    salida = BytesIO()
    documento = canvas.Canvas(salida, pagesize=(300, 300))
    for _ in range(cantidad):
        documento.drawImage(ImageReader(imagen), 0, 0, width=300, height=300)
        documento.showPage()
    documento.save()
    return Archivo("escaneado.pdf", salida.getvalue(), "application/pdf")


class Archivo(BytesIO):
    def __init__(self, nombre, contenido, tipo="text/plain"):
        super().__init__(contenido)
        self.name = nombre
        self.type = tipo
        self.size = len(contenido)


def quiz_ejemplo(cantidad=3):
    return Quiz.model_validate({"preguntas": [
        {"pregunta": f"Pregunta {numero}?", "opciones": ["A", "B", "C", "D"],
         "indice_correcta": 0, "explicacion": "La respuesta se encuentra en el documento."}
        for numero in range(cantidad)
    ]})


class DatosTemporales(unittest.TestCase):
    def setUp(self):
        temporal = tempfile.TemporaryDirectory()
        self.addCleanup(temporal.cleanup)
        self.carpeta = Path(temporal.name)
        for atributo, valor in (
            ("DATA_DIR", self.carpeta),
            ("CONVERSACIONES_PATH", self.carpeta / "conversaciones.json"),
            ("HISTORIAL_ANTERIOR_PATH", self.carpeta / "historial.json"),
        ):
            parche = patch.object(storage, atributo, valor)
            parche.start()
            self.addCleanup(parche.stop)
        self.contexto = {
            "documento_fragmentos": ["Archivo: agua.txt\nEl agua hierve a cien grados."],
            "documento_nombres": ["agua.txt"],
            "audio_transcripciones": "Consulta sobre agua.", "audio_nombres": ["audio.wav"],
            "imagenes_activas": [{"nombre": "agua.png", "tipo": "image/png", "bytes": b'imagen-local'}],
            "metadatos": [{"nombre": "agua.txt", "tamano": 50, "extension": "txt"}],
        }


class TestStorage(DatosTemporales):
    def test_recarga_contexto_y_bytes_fuera_del_json(self):
        storage.guardar_contexto("abc123", self.contexto)
        self.assertEqual(storage.cargar_contexto("abc123"), self.contexto)
        serializado = json.loads((self.carpeta / "contextos/abc123.json").read_text(encoding="utf-8"))
        self.assertNotIn("bytes", serializado["imagenes_activas"][0])

    def test_imagen_corrupta_y_faltante(self):
        storage.guardar_contexto("abc123", self.contexto)
        imagen = next((self.carpeta / "archivos/abc123").iterdir())
        imagen.write_bytes(b"corrupta")
        self.assertEqual(storage.cargar_contexto("abc123")["imagenes_activas"], [])
        imagen.unlink()
        recuperado = storage.cargar_contexto("abc123")
        self.assertEqual(recuperado["imagenes_activas"], [])
        self.assertEqual(recuperado["documento_fragmentos"], self.contexto["documento_fragmentos"])

    def test_json_corrupto_formas_invalidas_y_eliminacion(self):
        storage.guardar_contexto("abc123", self.contexto)
        ruta = self.carpeta / "contextos/abc123.json"
        for contenido in ("{invalido", "null", "[]", '{"documento_fragmentos": 5, "imagenes_activas": [null]}'):
            ruta.write_text(contenido, encoding="utf-8")
            self.assertFalse(storage.cargar_contexto("abc123").get("documento_fragmentos"))
        storage.eliminar_contexto("abc123")
        storage.eliminar_contexto("abc123")
        self.assertFalse(ruta.exists())
        self.assertFalse((self.carpeta / "archivos/abc123").exists())
        with self.assertRaises(ValueError):
            storage.cargar_contexto("../otra")

    def test_conversaciones_migracion_y_busqueda(self):
        storage.HISTORIAL_ANTERIOR_PATH.write_text(json.dumps([{"role": "user", "content": "El agua"}]), encoding="utf-8")
        chats = storage.cargar_conversaciones()
        self.assertEqual(chats[0]["titulo"], "Conversación anterior")
        self.assertEqual(storage.buscar_conversaciones(chats, "AGUA"), chats)
        self.assertIn("El agua", storage.obtener_fragmento_coincidente(chats[0], "agua"))
        self.assertEqual(storage.cargar_conversaciones(), chats)


class TestArchivos(unittest.TestCase):
    def test_pdf_escaneado_lectura_visual(self):
        pdf = pdf_escaneado()
        self.assertFalse(archivos.leer_pdf(pdf).strip())
        cliente = Mock()
        cliente.chat.completions.create.return_value = SimpleNamespace(choices=[SimpleNamespace(
            finish_reason="stop", message=SimpleNamespace(content="CODIGO DE ETICA\nTrabaja con honestidad. Respeta a tus companeros."),
        )])
        pdf.seek(0)
        avisos, errores, contexto = [], [], {}
        self.assertTrue(archivos.procesar_archivos_adjuntos([pdf], contexto, cliente, avisos.append, errores.append))
        self.assertEqual(errores, [])
        self.assertTrue(any("honestidad" in fragmento for fragmento in contexto["documento_fragmentos"]))
        self.assertEqual(contexto["metadatos"][0]["paginas_lectura_visual"], [1])
        solicitud = cliente.chat.completions.create.call_args.kwargs
        self.assertEqual(solicitud["model"], "gpt-4o-mini")
        datos = solicitud["messages"][1]["content"][0]["file"]["file_data"]
        enviado = PdfReader(BytesIO(base64.b64decode(datos.split(",", 1)[1])))
        self.assertEqual(len(enviado.pages), 1)
        self.assertTrue(enviado.pages[0].images)

    def test_lectura_visual_limita_paginas(self):
        cliente = Mock()
        cliente.chat.completions.create.return_value = SimpleNamespace(choices=[SimpleNamespace(
            finish_reason="stop", message=SimpleNamespace(content="Trabaja con honestidad."),
        )])
        avisos, errores, contexto = [], [], {}
        archivos.procesar_archivos_adjuntos(
            [pdf_escaneado(3)], contexto, cliente, avisos.append, errores.append, max_paginas_ocr=1,
        )
        self.assertEqual(cliente.chat.completions.create.call_count, 1)
        self.assertEqual(contexto["metadatos"][0]["paginas_omitidas"], [2, 3])
        self.assertTrue(any("revision sera parcial" in aviso for aviso in avisos))

    def test_pdf_mixto_conserva_texto_y_lectura_visual(self):
        from pypdf import PdfWriter

        documento = archivos.DocumentoGenerado(titulo="Reglas", secciones=[{"encabezado": "Respeto", "contenido": "Respeta a tus companeros."}])
        texto, _, _ = archivos.exportar_documento(documento, "PDF")
        salida = BytesIO()
        escritor = PdfWriter()
        escritor.add_page(PdfReader(BytesIO(texto)).pages[0])
        escritor.add_page(PdfReader(pdf_escaneado()).pages[0])
        escritor.write(salida)
        cliente = Mock()
        cliente.chat.completions.create.return_value = SimpleNamespace(choices=[SimpleNamespace(
            finish_reason="stop", message=SimpleNamespace(content="Trabaja con honestidad."),
        )])
        contexto = {}
        archivos.procesar_archivos_adjuntos([Archivo("mixto.pdf", salida.getvalue(), "application/pdf")], contexto, cliente, Mock(), Mock())
        contenido = " ".join(contexto["documento_fragmentos"])
        self.assertIn("Respeta a tus companeros", contenido)
        self.assertIn("Trabaja con honestidad", contenido)
        self.assertEqual(cliente.chat.completions.create.call_count, 1)
        self.assertEqual(contexto["metadatos"][0]["paginas_lectura_visual"], [2])

    def test_lectura_visual_fallida_no_bloquea_txt(self):
        cliente = Mock()
        cliente.chat.completions.create.side_effect = RuntimeError("detalle tecnico privado")
        errores, contexto = [], {}
        archivos.procesar_archivos_adjuntos(
            [pdf_escaneado(), Archivo("valido.txt", b"Texto legible")], contexto, cliente, Mock(), errores.append,
        )
        self.assertEqual(contexto["documento_nombres"], ["valido.txt"])
        self.assertTrue(errores)
        self.assertNotIn("detalle tecnico", errores[0])

    def test_limites_procesan_los_demas(self):
        avisos = []
        adjuntos = [Archivo("grande.txt", b"x" * 33), Archivo("clase.txt", b"El agua hierve.")]
        permitidos = archivos.filtrar_archivos(adjuntos, 5, 32, 20, avisos.append)
        self.assertEqual([archivo.name for archivo in permitidos], ["clase.txt"])
        contexto = {}
        self.assertTrue(archivos.procesar_archivos_adjuntos(permitidos, contexto, Mock(), avisos.append, avisos.append))
        self.assertEqual(contexto["documento_nombres"], ["clase.txt"])
        self.assertEqual(len(archivos.filtrar_archivos(adjuntos, 1, 100, 20, avisos.append)), 1)
        self.assertEqual(archivos.filtrar_archivos([Archivo("audio.wav", b"x" * 21)], 5, 100, 20, avisos.append), [])

    def test_error_de_un_archivo_no_interrumpe_los_demas(self):
        errores = []
        contexto = {}
        archivos.procesar_archivos_adjuntos(
            [Archivo("roto.pdf", b"no-es-pdf"), Archivo("valido.txt", b"Texto valido")],
            contexto, Mock(), Mock(), errores.append,
        )
        self.assertEqual(contexto["documento_nombres"], ["valido.txt"])
        self.assertEqual(len(errores), 1)
        self.assertNotIn("Traceback", errores[0])

    def test_lectores_docx_xlsx_y_busqueda(self):
        documento = Document()
        documento.add_paragraph("Agua potable")
        documento.add_table(rows=1, cols=1).cell(0, 0).text = "Cien grados"
        word = BytesIO()
        documento.save(word)
        texto = archivos.leer_archivo(Archivo("agua.docx", word.getvalue()))
        self.assertIn("Cien grados", texto)
        libro = Workbook()
        libro.active.append(["Agua", 100])
        excel = BytesIO()
        libro.save(excel)
        self.assertIn("Agua | 100", archivos.leer_archivo(Archivo("agua.xlsx", excel.getvalue())))
        fragmentos = dividir_texto(texto, tamano=2)
        self.assertTrue(buscar_informacion_relevante("agua", fragmentos))
        self.assertEqual(buscar_informacion_relevante("?", ["!"]), [])

    def test_transcripcion_con_cliente_simulado(self):
        cliente = Mock()
        cliente.audio.transcriptions.create.return_value = SimpleNamespace(text="Explica el agua")
        self.assertEqual(archivos.transcribir_audio(Archivo("microfono.wav", b"audio", "audio/wav"), cliente), "Explica el agua")
        self.assertEqual(cliente.audio.transcriptions.create.call_args.kwargs["model"], "gpt-4o-mini-transcribe")

    def test_errores_diferenciados_sin_detalles_tecnicos(self):
        peticion = httpx.Request("POST", "https://api.openai.com/v1/test")
        respuesta = httpx.Response(429, request=peticion)
        casos = [
            (RateLimitError("SECRETO", response=respuesta, body={"code": "insufficient_quota"}), "Saldo o cuota"),
            (RateLimitError("SECRETO", response=respuesta, body={"code": "rate_limit_exceeded"}), "temporalmente"),
            (AuthenticationError("SECRETO", response=respuesta, body={}), "clave"),
            (NotFoundError("SECRETO", response=respuesta, body={}), "modelo"),
            (APIConnectionError(request=peticion), "conectar"),
            (RuntimeError("SECRETO"), "Vuelve a intentarlo"),
        ]
        for error, esperado in casos:
            mensaje = archivos.mensaje_error(error)
            self.assertIn(esperado, mensaje)
            self.assertNotIn("SECRETO", mensaje)

    def test_documentos_exportables_y_generacion(self):
        documento = archivos.DocumentoGenerado(titulo="Resumen del agua", secciones=[
            {"encabezado": "Temperatura", "contenido": "El agua hierve a cien grados. Ácido & básico <texto>."},
        ])
        cliente = Mock()
        cliente.chat.completions.parse.return_value = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(parsed=documento))])
        self.assertEqual(archivos.generar_documento(cliente, "gpt-4o-mini", "Resumen", "Agua"), documento)
        texto, _, _ = archivos.exportar_documento(documento, "TXT")
        self.assertIn("Ácido", texto.decode("utf-8"))
        word, _, _ = archivos.exportar_documento(documento, "DOCX")
        self.assertEqual(Document(BytesIO(word)).paragraphs[0].text, documento.titulo)
        pdf, mime, extension = archivos.exportar_documento(documento, "PDF")
        self.assertEqual((mime, extension), ("application/pdf", "pdf"))
        self.assertIn("cien grados", PdfReader(BytesIO(pdf)).pages[0].extract_text())


class TestQuiz(unittest.TestCase):
    def test_validacion_nota_y_respuesta_incompleta(self):
        quiz = quiz_ejemplo()
        self.assertEqual(calificar_quiz(quiz, [0, 1, 0]), (2, 66.7))
        for campo, valor in (("indice_correcta", 4), ("indice_correcta", True), ("opciones", ["A"] * 4), ("explicacion", " ")):
            datos = quiz.model_dump()
            datos["preguntas"][0][campo] = valor
            with self.assertRaises(ValidationError):
                Quiz.model_validate(datos)
        with self.assertRaises(ValueError):
            calificar_quiz(quiz, [None, 0, 0])

    def test_generacion_cantidad_y_rechazo(self):
        cliente = Mock()
        mensaje = SimpleNamespace(parsed=quiz_ejemplo())
        cliente.chat.completions.parse.return_value = SimpleNamespace(choices=[SimpleNamespace(message=mensaje)])
        self.assertEqual(len(generar_quiz(cliente, "gpt-4o-mini", ["Agua"], 3, "Basico").preguntas), 3)
        for cantidad in (5, 11):
            with self.assertRaises(ValueError):
                generar_quiz(cliente, "gpt-4o-mini", ["Agua"], cantidad, "Basico")
        mensaje.parsed = None
        with self.assertRaises(ValueError):
            generar_quiz(cliente, "gpt-4o-mini", ["Agua"], 3, "Basico")


class TestInterfaz(DatosTemporales):
    def setUp(self):
        super().setUp()
        parche = patch.dict(os.environ, {"OPENAI_API_KEY": "clave-local-de-prueba"})
        parche.start()
        self.addCleanup(parche.stop)
        self.chats = [storage.crear_conversacion("Con documento"), storage.crear_conversacion("Sin documento")]
        storage.guardar_conversaciones(self.chats)
        storage.guardar_contexto(self.chats[0]["id"], self.contexto)

    def abrir(self):
        app = AppTest.from_file(str(ROOT / "chatbot_v5.py"), default_timeout=30).run()
        self.assertFalse(app.exception, [error.message for error in app.exception])
        return app

    def boton(self, app, etiqueta):
        return next(boton for boton in app.button if boton.label == etiqueta)

    def test_modulo_archivos_anterior_se_actualiza_antes_de_procesar(self):
        def procesador_anterior(adjuntos, contexto, cliente, advertir, reportar_error):
            raise AssertionError("No se debe invocar el procesador anterior.")

        entrada = SimpleNamespace(text="revisa este documento", files=[
            Archivo("legible.txt", b"Trabaja con honestidad y respeta a tus companeros."),
        ], audio=None)
        with patch.object(archivos, "procesar_archivos_adjuntos", procesador_anterior), patch("streamlit.chat_input", return_value=entrada), patch("langchain_openai.ChatOpenAI") as modelo:
            modelo.return_value.bind_tools.return_value = modelo.return_value
            modelo.return_value.stream.return_value = [AIMessageChunk(content="El texto promueve honestidad y respeto.")]
            app = self.abrir()
            self.assertFalse(app.exception)
            self.assertIsNot(archivos.procesar_archivos_adjuntos, procesador_anterior)
            self.assertIn("Trabaja con honestidad", modelo.return_value.stream.call_args.args[0][0].content)
        self.assertEqual(storage.cargar_contexto(self.chats[0]["id"])["documento_nombres"], ["legible.txt"])

    def test_recarga_cambio_quitar_y_eliminar(self):
        app = self.abrir()
        self.assertFalse(self.boton(app, "Generar quiz").disabled)
        self.boton(app, "Sin documento").click().run()
        self.assertTrue(self.boton(app, "Generar quiz").disabled)
        self.boton(app, "Con documento").click().run()
        self.assertFalse(self.boton(app, "Generar quiz").disabled)
        recargada = self.abrir()
        self.assertEqual(recargada.session_state["contextos_archivos_v5"][self.chats[0]["id"]], self.contexto)
        self.boton(recargada, "Quitar").click().run()
        self.assertFalse(recargada.exception)
        self.assertTrue(self.boton(recargada, "Generar quiz").disabled)
        self.assertFalse((self.carpeta / "contextos" / f"{self.chats[0]['id']}.json").exists())
        storage.guardar_contexto(self.chats[0]["id"], self.contexto)
        app = self.abrir()
        next(boton for boton in app.button if boton.key == f"eliminar_chat_{self.chats[0]['id']}").click().run()
        self.boton(app, "Eliminar").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.session_state["conversaciones"]), 1)
        self.assertFalse((self.carpeta / "archivos" / self.chats[0]["id"]).exists())

    def test_preguntar_tras_recarga_sin_adjuntar(self):
        app = self.abrir()
        with patch("langchain_openai.ChatOpenAI") as modelo:
            modelo.return_value.bind_tools.return_value = modelo.return_value
            modelo.return_value.stream.return_value = [SimpleNamespace(content="Hierve a cien grados.")]
            app.chat_input[0].set_value("A que temperatura hierve el agua?").run()
            self.assertFalse(app.exception)
            mensajes = modelo.return_value.stream.call_args.args[0]
            self.assertIn("El agua hierve a cien grados", mensajes[0].content)
            self.assertEqual(app.session_state["messages"][-1]["content"], "Hierve a cien grados.")
        self.assertEqual(storage.cargar_conversaciones()[0]["mensajes"][-1]["role"], "assistant")

    def test_revision_generica_documento_persistido_sin_coincidencias(self):
        self.contexto["documento_fragmentos"] = [
            "Archivo: Codigo de etica.pdf\nLa integridad exige honestidad y responsabilidad profesional.",
            "La confidencialidad protege la informacion privada.",
        ]
        self.contexto["documento_nombres"] = ["Codigo de etica.pdf"]
        self.assertEqual(buscar_informacion_relevante("revisa este documento", self.contexto["documento_fragmentos"]), [])
        storage.guardar_contexto(self.chats[0]["id"], self.contexto)
        app = self.abrir()
        with patch("langchain_openai.ChatOpenAI") as modelo:
            modelo.return_value.bind_tools.return_value = modelo.return_value
            modelo.return_value.stream.return_value = [AIMessageChunk(content="El texto destaca la integridad y la confidencialidad.")]
            app.chat_input[0].set_value("revisa este documento").run()
            self.assertFalse(app.exception)
            sistema = modelo.return_value.stream.call_args.args[0][0].content
            for fragmento in self.contexto["documento_fragmentos"]:
                self.assertIn(fragmento, sistema)
            self.assertIn("Sí puedes revisar estos documentos", sistema)

    def test_revision_generica_pdf_recien_adjuntado(self):
        documento = archivos.DocumentoGenerado(
            titulo="Codigo de etica", secciones=[{
                "encabezado": "Integridad", "contenido": "Honestidad, confidencialidad y responsabilidad profesional.",
            }],
        )
        contenido, _, _ = archivos.exportar_documento(documento, "PDF")
        pdf = Archivo("Codigo de etica.pdf", contenido, "application/pdf")
        entrada = SimpleNamespace(text="revisa este documento", files=[pdf], audio=None)
        with patch("streamlit.chat_input", return_value=entrada), patch("langchain_openai.ChatOpenAI") as modelo:
            modelo.return_value.bind_tools.return_value = modelo.return_value
            modelo.return_value.stream.return_value = [AIMessageChunk(content="El codigo aborda la integridad profesional.")]
            app = self.abrir()
            sistema = modelo.return_value.stream.call_args.args[0][0].content
            self.assertIn("Honestidad, confidencialidad y responsabilidad profesional", sistema)
            self.assertIn("Codigo de etica.pdf", sistema)
        recuperado = storage.cargar_contexto(self.chats[0]["id"])
        self.assertEqual(recuperado["documento_nombres"], ["Codigo de etica.pdf"])
        self.assertTrue(any("confidencialidad" in fragmento for fragmento in recuperado["documento_fragmentos"]))

    def test_pdf_escaneado_activa_contexto_y_quiz_tras_recarga(self):
        storage.eliminar_contexto(self.chats[0]["id"])
        entrada = SimpleNamespace(text="revisa este documento", files=[pdf_escaneado()], audio=None)
        cliente = Mock()
        cliente.chat.completions.create.return_value = SimpleNamespace(choices=[SimpleNamespace(
            finish_reason="stop", message=SimpleNamespace(content="CODIGO DE ETICA\nTrabaja con honestidad y respeta a tus companeros."),
        )])
        with patch("streamlit.chat_input", side_effect=[entrada, None]), patch("openai.OpenAI", return_value=cliente), patch("langchain_openai.ChatOpenAI") as modelo:
            modelo.return_value.bind_tools.return_value = modelo.return_value
            modelo.return_value.stream.return_value = [AIMessageChunk(content="El codigo promueve honestidad y respeto.")]
            app = self.abrir()
            sistema = modelo.return_value.stream.call_args.args[0][0].content
            self.assertIn("Trabaja con honestidad", sistema)
            app.run()
            self.assertFalse(app.exception)
            self.assertFalse(self.boton(app, "Generar quiz").disabled)
            self.assertEqual(cliente.chat.completions.create.call_count, 1)
        recuperado = storage.cargar_contexto(self.chats[0]["id"])
        self.assertEqual(recuperado["metadatos"][0]["paginas_lectura_visual"], [1])
        recargada = self.abrir()
        self.assertFalse(self.boton(recargada, "Generar quiz").disabled)

    def test_revision_generica_extensa_respeta_limite_contexto(self):
        self.contexto["documento_fragmentos"] = [
            f"SECCION_{numero} " + "integridad confidencialidad " * 1000 for numero in range(6)
        ]
        storage.guardar_contexto(self.chats[0]["id"], self.contexto)
        app = self.abrir()
        with patch("langchain_openai.ChatOpenAI") as modelo:
            modelo.return_value.bind_tools.return_value = modelo.return_value
            modelo.return_value.stream.return_value = [AIMessageChunk(content="Revision de los fragmentos disponibles.")]
            app.chat_input[0].set_value("revisa este documento").run()
            self.assertFalse(app.exception)
            sistema = modelo.return_value.stream.call_args.args[0][0].content
            for numero in range(6):
                self.assertIn(f"SECCION_{numero}", sistema)
            self.assertLess(len(sistema), 63000)
            self.assertTrue(any("material es extenso" in aviso.value for aviso in app.warning))

    def test_pdf_sin_texto_no_envia_consulta_al_modelo(self):
        from pypdf import PdfWriter

        salida = BytesIO()
        escritor = PdfWriter()
        escritor.add_blank_page(width=300, height=300)
        escritor.write(salida)
        pdf = Archivo("escaneado.pdf", salida.getvalue(), "application/pdf")
        entrada = SimpleNamespace(text="revisa este documento", files=[pdf], audio=None)
        with patch("streamlit.chat_input", return_value=entrada), patch("langchain_openai.ChatOpenAI") as modelo:
            app = self.abrir()
            modelo.assert_not_called()
        self.assertTrue(any("PDF escaneado" in aviso.value for aviso in app.warning))
        self.assertTrue(any("consulta no se envió" in aviso.value for aviso in app.warning))
        self.assertEqual(len(app.session_state["messages"]), 0)

    def test_adjunto_historico_sin_contexto_pide_recuperarlo(self):
        storage.eliminar_contexto(self.chats[0]["id"])
        self.chats[0]["mensajes"] = [
            {"role": "user", "content": "revisa este documento\n\n📎 Codigo de etica.pdf"},
            {"role": "assistant", "content": "No tengo la capacidad de revisar documentos directamente."},
        ]
        storage.guardar_conversaciones(self.chats)
        app = self.abrir()
        self.assertTrue(any("contenido no está activo" in aviso.value for aviso in app.warning))
        with patch("langchain_openai.ChatOpenAI") as modelo:
            modelo.return_value.bind_tools.return_value = modelo.return_value
            modelo.return_value.stream.return_value = [AIMessageChunk(content="Adjunta el documento de nuevo para revisarlo.")]
            app.chat_input[0].set_value("revisa este documento").run()
            sistema = modelo.return_value.stream.call_args.args[0][0].content
            self.assertIn("El nombre de un archivo en el historial no equivale a tener su contenido", sistema)
            self.assertIn("pide adjuntarlo de nuevo", sistema)

    def test_adjunto_fallido_no_bloquea_otro_documento_valido(self):
        entrada = SimpleNamespace(text="revisa estos documentos", files=[
            Archivo("roto.pdf", b"pdf-invalido", "application/pdf"),
            Archivo("legible.txt", b"Honestidad y responsabilidad profesional."),
        ], audio=None)
        with patch("streamlit.chat_input", return_value=entrada), patch("langchain_openai.ChatOpenAI") as modelo:
            modelo.return_value.bind_tools.return_value = modelo.return_value
            modelo.return_value.stream.return_value = [AIMessageChunk(content="Puedo revisar el TXT legible.")]
            app = self.abrir()
            sistema = modelo.return_value.stream.call_args.args[0][0].content
            self.assertIn("Honestidad y responsabilidad profesional", sistema)
            self.assertTrue(app.error)
        self.assertEqual(storage.cargar_contexto(self.chats[0]["id"])["documento_nombres"], ["legible.txt"])

    def test_quiz_formulario_y_explicaciones(self):
        app = self.abrir()
        with patch("quiz.generar_quiz", return_value=quiz_ejemplo(5)):
            self.boton(app, "Generar quiz").click().run()
        self.assertFalse(app.exception)
        self.boton(app, "Calificar").click().run()
        self.assertTrue(any("Responde todas" in aviso.value for aviso in app.warning))
        for radio in app.radio:
            radio.set_value(0)
        self.boton(app, "Calificar").click().run()
        self.assertTrue(any("100/100" in mensaje.value for mensaje in app.success))
        self.assertFalse(app.exception)
        self.boton(app, "Sin documento").click().run()
        self.assertEqual(len(app.radio), 0)

    def test_error_de_chat_no_se_guarda_como_respuesta(self):
        app = self.abrir()
        with patch("langchain_openai.ChatOpenAI") as modelo:
            modelo.return_value.bind_tools.return_value = modelo.return_value
            modelo.return_value.stream.side_effect = RuntimeError("detalle-tecnico-no-visible")
            app.chat_input[0].set_value("Explica el agua").run()
        self.assertFalse(app.exception)
        self.assertTrue(app.error)
        self.assertNotIn("detalle-tecnico", app.error[0].value)
        self.assertEqual(app.session_state["messages"][-1]["role"], "user")

    def test_recarga_con_chat_seleccionado_en_url(self):
        app = AppTest.from_file(str(ROOT / "chatbot_v5.py"), default_timeout=30)
        app.query_params["chat"] = self.chats[1]["id"]
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual(app.session_state["chat_activo"], self.chats[1]["id"])
        self.assertTrue(self.boton(app, "Generar quiz").disabled)

    def test_microfono_envia_consulta_y_persiste_transcripcion(self):
        grabacion = Archivo("microfono.wav", b"audio-simulado", "audio/wav")
        entrada = SimpleNamespace(text="", files=[], audio=grabacion)
        with patch("streamlit.chat_input", return_value=entrada) as campo, patch("archivos.transcribir_audio", return_value="Explica el agua") as transcripcion, patch("langchain_openai.ChatOpenAI") as modelo:
            modelo.return_value.bind_tools.return_value = modelo.return_value
            modelo.return_value.stream.return_value = [SimpleNamespace(content="El agua es esencial.")]
            app = self.abrir()
            self.assertTrue(campo.call_args.kwargs["accept_audio"])
            transcripcion.assert_called_once()
            self.assertEqual(app.session_state["messages"][-2]["content"], "Explica el agua")
            self.assertEqual(app.session_state["messages"][-1]["content"], "El agua es esencial.")
            self.assertFalse(any(boton.label == "Enviar grabación" for boton in app.button))
        contexto = storage.cargar_contexto(self.chats[0]["id"])
        self.assertIn("Explica el agua", contexto["audio_transcripciones"])
        self.assertEqual(contexto["metadatos"][-1]["origen"], "microfono")

    def test_generacion_y_descarga_documento_en_interfaz(self):
        app = self.abrir()
        documento = archivos.DocumentoGenerado(
            titulo="Resumen del agua", secciones=[{"encabezado": "Temperatura", "contenido": "Cien grados."}],
        )
        argumentos = json.dumps({"formato": "PDF", "instrucciones": "Prepara un resumen del agua"})
        llamadas = [
            AIMessageChunk(content="", tool_call_chunks=[{
                "name": "generar_documento", "args": argumentos[:25], "id": "solicitud_pdf", "index": 0,
            }]),
            AIMessageChunk(content="", tool_call_chunks=[{
                "name": None, "args": argumentos[25:], "id": None, "index": 0,
            }]),
        ]
        with patch("archivos.generar_documento", return_value=documento) as generador, patch("langchain_openai.ChatOpenAI") as modelo:
            modelo.return_value.bind_tools.return_value = modelo.return_value
            modelo.return_value.stream.return_value = llamadas
            app.chat_input[0].set_value("Genera un PDF del agua").run()
            generador.assert_called_once()
            self.assertIn("El agua hierve", generador.call_args.args[3])
        self.assertFalse(app.exception)
        self.assertFalse(any(boton.label == "Generar documento" for boton in app.button))
        self.assertTrue(any(encabezado.value == "Temperatura" for encabezado in app.subheader))
        self.assertEqual(app.session_state["messages"][-1]["documento"]["formato"], "PDF")
        recargada = self.abrir()
        self.assertTrue(any(encabezado.value == "Temperatura" for encabezado in recargada.subheader))
        self.assertEqual(storage.cargar_conversaciones()[0]["mensajes"][-1]["documento"]["formato"], "PDF")
        with patch("langchain_openai.ChatOpenAI") as modelo:
            modelo.return_value.bind_tools.return_value = modelo.return_value
            modelo.return_value.stream.return_value = [AIMessageChunk(content="El documento trata del agua.")]
            recargada.chat_input[0].set_value("De que trata el documento que generaste?").run()
            mensajes = modelo.return_value.stream.call_args.args[0]
            self.assertTrue(any("Cien grados" in mensaje.content for mensaje in mensajes[1:]))
        self.boton(app, "Sin documento").click().run()
        self.assertFalse(any(encabezado.value == "Temperatura" for encabezado in app.subheader))

    def test_word_y_txt_solicitados_desde_el_chat(self):
        documento = archivos.DocumentoGenerado(
            titulo="Agua", secciones=[{"encabezado": "Resumen", "contenido": "El agua hierve."}],
        )
        for formato, consulta in (("DOCX", "Exporta ese resumen a Word"), ("TXT", "Guardalo en un TXT")):
            with self.subTest(formato=formato):
                app = self.abrir()
                llamada = AIMessageChunk(content="", tool_call_chunks=[{
                    "name": "generar_documento", "args": json.dumps({"formato": formato, "instrucciones": consulta}),
                    "id": "exportar", "index": 0,
                }])
                with patch("archivos.generar_documento", return_value=documento), patch("langchain_openai.ChatOpenAI") as modelo:
                    modelo.return_value.bind_tools.return_value = modelo.return_value
                    modelo.return_value.stream.return_value = [llamada]
                    app.chat_input[0].set_value(consulta).run()
                self.assertFalse(app.exception)
                self.assertEqual(app.session_state["messages"][-1]["documento"]["formato"], formato)

    def test_pregunta_sobre_pdf_sigue_siendo_chat(self):
        app = self.abrir()
        with patch("archivos.generar_documento") as generador, patch("langchain_openai.ChatOpenAI") as modelo:
            modelo.return_value.bind_tools.return_value = modelo.return_value
            modelo.return_value.stream.return_value = [AIMessageChunk(content="PDF es un formato de documento.")]
            app.chat_input[0].set_value("Que es un PDF?").run()
            generador.assert_not_called()
        self.assertFalse(app.exception)
        self.assertNotIn("documento", app.session_state["messages"][-1])

    def test_herramienta_o_formato_invalido_no_genera_archivo(self):
        for nombre, formato in (("ejecutar_codigo", "PDF"), ("generar_documento", "EXE")):
            with self.subTest(nombre=nombre, formato=formato):
                app = self.abrir()
                llamada = AIMessageChunk(content="", tool_call_chunks=[{
                    "name": nombre, "args": json.dumps({"formato": formato, "instrucciones": "Solicitud"}),
                    "id": "invalida", "index": 0,
                }])
                with patch("archivos.generar_documento") as generador, patch("langchain_openai.ChatOpenAI") as modelo:
                    modelo.return_value.bind_tools.return_value = modelo.return_value
                    modelo.return_value.stream.return_value = [llamada]
                    app.chat_input[0].set_value("Genera un documento").run()
                    generador.assert_not_called()
                self.assertFalse(app.exception)
                self.assertTrue(app.error)
                self.assertEqual(app.session_state["messages"][-1]["role"], "user")


if __name__ == "__main__":
    unittest.main()